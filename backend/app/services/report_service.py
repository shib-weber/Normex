import json
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document

def make_pdf(result):
    bio = BytesIO()
    c = canvas.Canvas(bio, pagesize=A4)
    w, h = A4
    y = h - 45
    c.setFont("Helvetica-Bold", 18)
    c.drawString(40, y, "NORMEX Procurement Intelligence Report")
    y -= 30
    c.setFont("Helvetica", 9)
    lines = [
        f"Product: {result['extraction']['product']}",
        f"Domain: {result['extraction']['domain']}",
        f"Coverage: {result['coverage']['covered']}/{result['coverage']['total']}",
        "",
        "Recommendations:",
    ]
    for r in result["recommendations"]:
        lines += [f"{r['category']} | {r['standard_number']} | {r['title']} | relevance {r['relevance']}%"]
    lines += ["", "Disclaimer:", result["disclaimer"]]
    for line in lines:
        if y < 45:
            c.showPage(); y = h - 45
        c.drawString(40, y, line[:125])
        y -= 14
    c.save()
    bio.seek(0)
    return bio

def make_docx(result):
    bio = BytesIO()
    doc = Document()
    doc.add_heading("NORMEX Procurement Intelligence Report", 0)
    doc.add_paragraph(f"Product: {result['extraction']['product']}")
    doc.add_paragraph(f"Domain: {result['extraction']['domain']}")
    doc.add_heading("Recommendations", level=1)
    for r in result["recommendations"]:
        doc.add_paragraph(f"{r['category']} — {r['standard_number']} — {r['title']} — relevance {r['relevance']}%")
    doc.add_heading("Requirement Coverage", level=1)
    doc.add_paragraph(f"{result['coverage']['covered']} / {result['coverage']['total']}")
    doc.add_heading("Disclaimer", level=1)
    doc.add_paragraph(result["disclaimer"])
    doc.save(bio)
    bio.seek(0)
    return bio
