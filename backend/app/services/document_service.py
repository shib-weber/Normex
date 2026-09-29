from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

from app.core.config import settings


ALLOWED = {".pdf", ".docx", ".txt"}


def extract_pdf_text(path: str) -> str:
    """
    Extract text from a PDF document.

    Uses PyMuPDF because it generally provides better PDF text extraction
    than many lightweight PDF parsers.
    """
    try:
        pages = []

        with fitz.open(path) as doc:
            for page in doc:
                text = page.get_text("text") or ""

                if text.strip():
                    pages.append(text.strip())

        return "\n\n".join(pages).strip()

    except Exception as exc:
        raise ValueError(f"Unable to read PDF document: {exc}") from exc


def extract_docx_text(path: str) -> str:
    """
    Extract text from DOCX paragraphs and tables.

    Tables are important because procurement documents and tenders
    frequently store specifications, quantities and requirements inside them.
    """
    try:
        doc = Document(path)

        sections = []

        # Normal paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            if text:
                sections.append(text)

        # Tables
        for table in doc.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    sections.append(" | ".join(cells))

        return "\n".join(sections).strip()

    except Exception as exc:
        raise ValueError(f"Unable to read DOCX document: {exc}") from exc


def extract_txt_text(path: str) -> str:
    """
    Extract text from a plain-text file.
    """
    try:
        return Path(path).read_text(
            encoding="utf-8",
            errors="ignore"
        ).strip()

    except Exception as exc:
        raise ValueError(f"Unable to read TXT document: {exc}") from exc


def extract_text(path: str) -> str:
    """
    Main document extraction function.

    Supported:
        - PDF
        - DOCX
        - TXT

    Returns:
        Extracted document text as a string.
    """

    p = Path(path)

    if not p.exists():
        raise FileNotFoundError(f"Document not found: {path}")

    if not p.is_file():
        raise ValueError(f"Path is not a file: {path}")

    extension = p.suffix.lower()

    if extension not in ALLOWED:
        raise ValueError(
            f"Unsupported document type '{extension}'. "
            f"Supported formats: {', '.join(sorted(ALLOWED))}"
        )

    if extension == ".pdf":
        text = extract_pdf_text(str(p))

    elif extension == ".docx":
        text = extract_docx_text(str(p))

    elif extension == ".txt":
        text = extract_txt_text(str(p))

    else:
        raise ValueError("Unsupported document type")

    if not text.strip():
        raise ValueError(
            "No readable text was found in the document. "
            "If this is a scanned PDF, OCR may be required."
        )

    return text.strip()


def extract_document_text(
    data: bytes,
    filename: str,
    content_type: str | None = None
) -> str:
    """
    Compatibility helper for routes that receive uploaded files
    as bytes instead of an already-saved file path.

    Supports PDF, DOCX and TXT.
    """

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED:
        raise ValueError(
            f"Unsupported document type '{extension}'. "
            f"Supported formats: {', '.join(sorted(ALLOWED))}"
        )

    # Write the uploaded bytes to a temporary file because the
    # existing extract_text() function works with file paths.
    import tempfile

    suffix = extension

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=suffix
    ) as temp_file:

        temp_file.write(data)
        temp_path = temp_file.name

    try:
        return extract_text(temp_path)

    finally:
        Path(temp_path).unlink(missing_ok=True)
